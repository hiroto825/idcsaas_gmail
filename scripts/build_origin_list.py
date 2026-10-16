"""既存企業発・大学発・研究所発に絞ったリストを出力する。

usage: python3 scripts/build_origin_list.py <all_items.json> <cls_dir> <results2_dir> <output_dir>
"""
import csv
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TARGET = ["既存企業発", "大学発", "研究所発"]
EXCLUDED = ["一般スタートアップ", "対象外"]
COLUMNS = [
    ("No", 6), ("掲載日", 12), ("サイト", 20), ("出自", 11), ("親組織", 24), ("判定", 12),
    ("企業・組織名", 28), ("タイトル", 60), ("概要", 60), ("日付確度", 10), ("URL", 50),
]
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def norm(url):
    return (url or "").strip().rstrip("/").lower()


def load(all_items_path, cls_dir, results2_dir):
    items = json.loads(Path(all_items_path).read_text(encoding="utf-8"))
    cls = {}
    for f in Path(cls_dir).glob("*.json"):
        for c in json.loads(f.read_text(encoding="utf-8")):
            cls[c["id"]] = c
    for it in items:
        c = cls.get(it["id"], {})
        it["origin"] = c.get("origin", "一般スタートアップ")
        it["parent_org"] = c.get("parent_org", "")
    seen = {(norm(i.get("url")) or i.get("title", ""), i.get("company", "")) for i in items}
    seen_url = {norm(i.get("url")) for i in items if i.get("url")}
    for f in sorted(Path(results2_dir).glob("*.json")):
        for it in json.loads(f.read_text(encoding="utf-8")).get("items", []):
            key = (norm(it.get("url")) or it.get("title", ""), it.get("company", ""))
            if key in seen or (it.get("url") and norm(it["url"]) in seen_url):
                continue
            seen.add(key)
            items.append(it)
    return items


def style_header(ws):
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    ws.freeze_panes = "A2"


def write_list(ws, rows):
    ws.append([c for c, _ in COLUMNS])
    for n, it in enumerate(rows, 1):
        ws.append([
            n, it.get("date", ""), it.get("site", ""), it.get("origin", ""), it.get("parent_org", ""),
            it.get("category", ""), it.get("company", ""), it.get("title", ""), it.get("summary_ja", ""),
            it.get("date_confidence", ""), it.get("url", ""),
        ])
        if it.get("url"):
            cell = ws.cell(row=n + 1, column=len(COLUMNS))
            cell.hyperlink = it["url"]
            cell.font = Font(color="0563C1", underline="single")
    for idx, (_, w) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    for r in ws.iter_rows(min_row=2):
        for cell in r[7:9]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    style_header(ws)
    ws.auto_filter.ref = ws.dimensions


def main(all_items_path, cls_dir, results2_dir, output_dir):
    items = load(all_items_path, cls_dir, results2_dir)
    in_period = [i for i in items if i.get("date_confidence") != "対象期間外(参考)"]
    target = [i for i in in_period if i["origin"] in TARGET]
    excluded = [i for i in in_period if i["origin"] in EXCLUDED]
    target.sort(key=lambda i: (TARGET.index(i["origin"]), i.get("date", ""), i.get("site", "")))

    sites = []
    for i in in_period:
        if i["site"] not in sites:
            sites.append(i["site"])

    wb = Workbook()
    ws = wb.active
    ws.title = "サイト別集計"
    ws.append(["サイト"] + TARGET + ["対象計", "除外(一般スタートアップ)", "除外(対象外)"])
    rows = []
    for s in sites:
        c = Counter(i["origin"] for i in in_period if i["site"] == s)
        rows.append([s] + [c[o] for o in TARGET] + [sum(c[o] for o in TARGET), c["一般スタートアップ"], c["対象外"]])
    rows.sort(key=lambda r: -r[4])
    for r in rows:
        ws.append(r)
    ws.append(["合計"] + [sum(r[k] for r in rows) for k in range(1, 7)])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 30
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 14
    style_header(ws)

    write_list(wb.create_sheet("対象リスト"), target)
    write_list(wb.create_sheet("除外"), sorted(excluded, key=lambda i: (i["origin"], i.get("site", ""))))

    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    wb.save(out / "origin_new_business_2026-09.xlsx")
    with open(out / "origin_new_business_2026-09.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow([c for c, _ in COLUMNS])
        for n, it in enumerate(target, 1):
            w.writerow([n, it.get("date", ""), it.get("site", ""), it["origin"], it.get("parent_org", ""),
                        it.get("category", ""), it.get("company", ""), it.get("title", ""),
                        it.get("summary_ja", ""), it.get("date_confidence", ""), it.get("url", "")])

    print(f"target={len(target)} excluded={len(excluded)}")
    for r in rows:
        print(r)


if __name__ == "__main__":
    main(*sys.argv[1:5])
