"""企業発 新規事業事例の候補洗い出し結果（バッチ別JSON）を1つのExcelにまとめる。

usage: python3 scripts/build_new_biz_cases.py <raw_dir> <output_xlsx>
"""
import json
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")

COMPANY_COLS = [
    ("No.", 5), ("上場企業名", 24), ("証券コード", 9), ("業種", 12), ("公式サイト", 30), ("調査状況", 9),
    ("確認した範囲", 50), ("候補件数", 8), ("収録件数", 8), ("保留件数", 8), ("除外件数", 8), ("残る確認事項", 45),
]
CASE_COLS = [
    ("No.", 5), ("事業主体の企業名", 22), ("親会社名・証券コード", 22), ("自社／子会社／スピンアウト", 12), ("事業名", 26),
    ("国・地域", 6), ("提供開始時期", 11), ("対象顧客・課題", 36), ("提供内容・価値", 40), ("収益モデル", 26),
    ("従来の事業・活用資産・新たな展開領域", 40), ("有料提供の根拠", 36), ("A：成功実績（暫定）", 40),
    ("B：独自性（暫定）", 50), ("判定", 7), ("保留・除外理由", 36), ("追加確認事項", 40), ("現在の提供状況", 10),
    ("対象期間の確認", 9), ("根拠URL・発表日", 60), ("確認日", 11),
]


def style(ws, cols, wrap_from=None):
    for idx, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = WRAP
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def fmt_a(a):
    a = a or {}
    parts = [a.get("provisional", "未確認")]
    if a.get("aspects"):
        parts.append("観点：" + "、".join(a["aspects"]))
    if a.get("facts"):
        parts.append(a["facts"])
    return "\n".join(parts)


def fmt_b(b):
    b = b or {}
    parts = [b.get("provisional", "未確認")]
    if b.get("aspects"):
        parts.append("観点：" + "、".join(b["aspects"]))
    for label, key in (("比較対象", "comparison"), ("違い", "difference"), ("顧客への効果", "customer_effect")):
        if b.get(key):
            parts.append(f"{label}：{b[key]}")
    return "\n".join(parts)


def fmt_sources(sources):
    return "\n".join(
        f"[{s.get('date') or '日付不明'}{'' if s.get('is_2026') else '・補足'}] {s.get('fact', '')} {s.get('url', '')}"
        for s in sources or []
    )


def load(raw_dir):
    batches = [json.loads(f.read_text(encoding="utf-8")) for f in sorted(Path(raw_dir).glob("batch*.json"))]
    companies, cases, seen = [], [], set()
    for b in batches:
        companies.extend(b.get("companies", []))
        for c in b.get("candidates", []):
            key = (c.get("parent") or c.get("company"), c.get("business_name"))
            if key in seen:
                continue
            seen.add(key)
            c["checked_date"] = (b.get("finished_at") or "")[:10]
            cases.append(c)
    return batches, companies, cases


def write_summary(ws, batches, companies, cases):
    judg = Counter(c.get("judgement") for c in cases)
    a_ok = lambda c: (c.get("a_success") or {}).get("provisional") == "該当見込み"
    b_ok = lambda c: (c.get("b_uniqueness") or {}).get("provisional") == "該当見込み"
    held = [c for c in cases if c.get("judgement") == "保留"]
    status = Counter(c.get("status") for c in companies)
    searches = sum(b.get("search_count") or 0 for b in batches)
    rows = [
        ("調査基準日", datetime.now().strftime("%Y-%m-%d")),
        ("対象期間", "2026-01-01〜2026-10-30（調査実行日まで）"),
        ("調査段階", "候補洗い出し（検索結果のスニペットのみ）。ネットワーク制限で企業サイト等の本文を閲覧できず、全件を『保留（本文未確認）』とした"),
        ("対象企業リストの取得元・基準日", "未取得（JPXの上場会社一覧に接続できず）。今回は6業種48社を調査者が選定"),
        ("把握した対象企業グループ数", "未把握（全上場企業のリスト未取得）"),
        ("調査済み／一部調査／未調査", f"調査済み 0 ／ 一部調査 {status.get('一部調査', 0)} ／ 未調査 上記48社以外の全上場企業"),
        ("重複排除後の候補件数", len(cases)),
        ("収録／保留／除外", f"収録 0 ／ 保留 {judg.get('保留', 0)} ／ 除外 {judg.get('除外', 0)}"),
        ("保留の暫定内訳（A・B該当見込み）",
         f"Aのみ {sum(a_ok(c) and not b_ok(c) for c in held)} ／ Bのみ {sum(b_ok(c) and not a_ok(c) for c in held)} ／ "
         f"A・B両方 {sum(a_ok(c) and b_ok(c) for c in held)} ／ いずれも未確認 {sum(not a_ok(c) and not b_ok(c) for c in held)}"),
        ("主な保留理由", "本文未確認（全件）。加えて有料提供の根拠、事業単位の数値、対象期間の確認が不足"),
        ("検索回数（実測）", searches),
        ("費用", "計測不可"),
        ("作業時間", "バッチごとの開始・終了時刻を『企業別の調査状況』の元データに記録。全社調査の実績ではない"),
        ("既存DBとの重複確認", "未実施（既存DBの提供なし）"),
    ]
    ws.append(["項目", "内容"])
    for r in rows:
        ws.append(list(r))
    style(ws, [("項目", 32), ("内容", 100)])


def main(raw_dir, output):
    batches, companies, cases = load(raw_dir)
    wb = Workbook()
    write_summary(wb.active, batches, companies, cases)
    wb.active.title = "①調査サマリー"

    ws = wb.create_sheet("②企業別の調査状況")
    ws.append([c for c, _ in COMPANY_COLS])
    for n, co in enumerate(companies, 1):
        mine = [c for c in cases if co["company"] in (c.get("parent") or "", c.get("company") or "")]
        ws.append([n, co.get("company"), co.get("code"), co.get("industry"), co.get("website"), co.get("status"),
                   co.get("checked_scope"), co.get("candidates"), 0,
                   sum(c.get("judgement") == "保留" for c in mine), co.get("excluded"), co.get("remaining")])
    style(ws, COMPANY_COLS)

    for title, judgement in (("③候補一覧（保留）", "保留"), ("④除外一覧", "除外")):
        ws = wb.create_sheet(title)
        ws.append([c for c, _ in CASE_COLS])
        for n, c in enumerate([c for c in cases if c.get("judgement") == judgement], 1):
            ws.append([
                n, c.get("company"), f"{c.get('parent', '')}（{c.get('code', '')}）", c.get("entity_type"),
                c.get("business_name"), c.get("country", "JP"), c.get("launch_date"), c.get("customer_problem"),
                c.get("offering_value"), c.get("revenue_model"), c.get("origin_assets_domain"), c.get("paid_evidence"),
                fmt_a(c.get("a_success")), fmt_b(c.get("b_uniqueness")), c.get("judgement"), c.get("reason"),
                c.get("followups"), c.get("current_status"), "確認済" if c.get("period_confirmed") else "要確認",
                fmt_sources(c.get("sources")), c.get("checked_date"),
            ])
        style(ws, CASE_COLS)

    ws = wb.create_sheet("⑤調査・収集の実績")
    ws.append(["検証観点", "指標", "値", "計算方法・注記"])
    required = ["business_name", "customer_problem", "offering_value", "paid_evidence", "launch_date", "sources"]
    held = [c for c in cases if c.get("judgement") == "保留"]
    filled = sum(1 for c in held for k in required if c.get(k) and "未確認" not in str(c.get(k)))
    total = len(held) * len(required)
    minutes = 0.0
    for b in batches:
        try:
            minutes += (datetime.fromisoformat(b["finished_at"]) - datetime.fromisoformat(b["started_at"])).total_seconds() / 60
        except (KeyError, TypeError, ValueError):
            pass
    for r in (
        ("収集件数", "総収録件数", 0, "本文未確認のため収録確定なし"),
        ("収集件数", "調査済み企業グループ当たりの収録件数", "算出不可", "収録0件・調査済み0社"),
        ("収集件数", "候補（保留）件数／一部調査企業数", f"{len(held)}／{len(companies)}", "参考値"),
        ("内容の充実度", "必須項目の充足率（保留候補）", f"{filled}/{total}" + (f"（{filled / total:.0%}）" if total else ""),
         "事業名・対象顧客・提供内容・有料根拠・提供開始時期・出典の6項目が埋まり『未確認』でない数／候補数×6"),
        ("内容の充実度", "根拠本文を確認できた割合", f"0/{len(cases)}（0%）", "ネットワーク制限で本文閲覧不可"),
        ("内容の充実度", "主な不足情報", "有料提供の根拠、事業単位の数値、提供開始時期", ""),
        ("コスト", "確認できる実費・収録1件当たりの費用", "計測不可", "API利用料等を取得できない／収録0件"),
        ("時間", "実測できる調査時間", f"約{minutes:.0f}分（6バッチの合計、並列実行）", "各バッチの開始・終了時刻の差の合計。全社調査の実績ではない"),
        ("時間", "収録1件当たりの作業時間", "算出不可", "収録0件"),
    ):
        ws.append(list(r))
    style(ws, [("検証観点", 14), ("指標", 34), ("値", 30), ("計算方法・注記", 70)])

    ws = wb.create_sheet("⑥引き継ぎ")
    ws.append(["区分", "内容"])
    done = "、".join(co["company"] for co in companies)
    for r in (
        ("調査済み", "なし（本文確認まで終えた企業はない）"),
        ("一部調査（スニペット段階）", done),
        ("未調査", "上記以外の全上場企業。企業リスト（JPX上場会社一覧）が未取得"),
        ("次にやること", "①ネットワークを開放し、JPX上場会社一覧で企業リストを作る ②③の候補の本文（商品・料金ページ、IR資料）を確認して収録／除外を確定 ③残りの業種へ順に拡大"),
        ("重複確認に使う一覧", "③候補一覧・④除外一覧（親会社名＋事業名で照合）"),
        ("元データ", "data/new_biz_cases/raw/batch*.json"),
    ):
        ws.append(list(r))
    style(ws, [("区分", 28), ("内容", 110)])

    Path(output).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(f"companies={len(companies)} cases={len(cases)}", dict(Counter(c.get("judgement") for c in cases)))


if __name__ == "__main__":
    main(*sys.argv[1:3])
