"""収集結果(JSON)を統合し、Excel / CSV のリストを出力する。

usage: python3 scripts/build_list.py <results_dir> <output_dir>
"""
import csv
import json
import sys
from collections import Counter, OrderedDict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SITES = [
    "PR TIMES",
    "企業IR・適時開示",
    "企業・大学・研究所の自社サイト",
    "会社四季報オンライン",
    "MONOist",
    "キーマンズネット",
    "JST START",
    "JST A-STEP 成果集",
    "経産省 大学発ベンチャーDB",
    "GlobeNewswire",
    "PR Newswire",
    "Fraunhofer Investment Hub",
    # 追加サイト（量の担保）
    "Business Wire",
    "@Press",
    "ValuePress",
    "THE BRIDGE",
    "EU-Startups",
]
ADDED_SITES = {"Business Wire", "@Press", "ValuePress", "THE BRIDGE", "EU-Startups"}
CATEGORIES = ["新規事業", "スタートアップ", "資金調達", "商用リリース"]
COLUMNS = [
    ("No", 6),
    ("掲載日", 12),
    ("サイト", 22),
    ("判定", 12),
    ("組織区分", 12),
    ("企業・組織名", 28),
    ("タイトル", 60),
    ("概要", 60),
    ("該当キーワード", 18),
    ("日付確度", 14),
    ("URL", 50),
    ("他サイト重複", 20),
]


def company_key(name):
    return name.split("(")[0].split("（")[0].strip().lower()


def normalize_url(url):
    return (url or "").strip().rstrip("/").lower()


SITE_ALIASES = {"適時開示": "企業IR・適時開示"}


def load(results_dir):
    items, queries, notes = [], Counter(), {}
    for path in sorted(Path(results_dir).glob("*.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        q = data.get("queries_run", 0)
        n = data.get("notes", "")
        site_default = data["items"][0]["site"] if data.get("items") else data.get("source_group")
        if isinstance(q, dict):
            queries.update({SITE_ALIASES.get(k, k): v for k, v in q.items()})
        else:
            queries[site_default] += q
        if isinstance(n, dict):
            for k, v in n.items():
                notes.setdefault(SITE_ALIASES.get(k, k), []).append(v)
        elif n:
            notes.setdefault(site_default, []).append(n)
        items.extend(data.get("items", []))
    return items, queries, notes


def dedupe(items):
    seen, out = set(), []
    for it in items:
        # まとめ記事のURLを複数社で共有するケースがあるため企業名も含める
        key = (normalize_url(it.get("url")) or it.get("title", ""), it.get("company", ""))
        if key in seen:
            continue
        seen.add(key)
        out.append(it)
    return out


def main(results_dir, output_dir):
    items, queries, notes = load(results_dir)
    items = dedupe(items)
    in_period = [i for i in items if i.get("date_confidence") != "対象期間外(参考)"]
    reference = [i for i in items if i.get("date_confidence") == "対象期間外(参考)"]
    site_order = {s: n for n, s in enumerate(SITES)}
    in_period.sort(key=lambda i: (site_order.get(i["site"], 99), i.get("date", ""), i.get("company", "")))

    sites_by_company = {}
    for it in in_period:
        sites_by_company.setdefault(company_key(it.get("company", "")), set()).add(it["site"])
    for it in in_period:
        others = sites_by_company.get(company_key(it.get("company", "")), set()) - {it["site"]}
        it["_dup"] = "、".join(sorted(others))

    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)

    def row(n, it):
        return [
            n, it.get("date", ""), it.get("site", ""), it.get("category", ""), it.get("org_type", ""),
            it.get("company", ""), it.get("title", ""), it.get("summary_ja", ""),
            it.get("matched_keyword", ""), it.get("date_confidence", ""), it.get("url", ""), it.get("_dup", ""),
        ]

    with open(out / "new_business_2026-09.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow([c for c, _ in COLUMNS])
        for n, it in enumerate(in_period, 1):
            w.writerow(row(n, it))

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")

    def write_sheet(ws, rows):
        ws.append([c for c, _ in COLUMNS])
        for n, it in enumerate(rows, 1):
            ws.append(row(n, it))
            url = it.get("url")
            if url:
                cell = ws.cell(row=n + 1, column=11)
                cell.hyperlink = url
                cell.font = Font(color="0563C1", underline="single")
        for idx, (_, width) in enumerate(COLUMNS, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        for cell in ws[1]:
            cell.fill, cell.font = header_fill, header_font
        for r in ws.iter_rows(min_row=2):
            for cell in r[6:8]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    # サマリー
    ws = wb.active
    ws.title = "サイト別集計"
    ws.append(["区分", "サイト", "検索クエリ数"] + CATEGORIES + ["合計(9月)", "うち日付確認済", "参考(期間外)", "備考"])
    counts = OrderedDict((s, Counter()) for s in SITES)
    for it in in_period:
        c = counts.setdefault(it["site"], Counter())
        c[it.get("category", "")] += 1
        c["_total"] += 1
        if it.get("date_confidence") == "確認済":
            c["_confirmed"] += 1
    for it in reference:
        counts.setdefault(it["site"], Counter())["_ref"] += 1
    for site, c in counts.items():
        ws.append(
            ["追加" if site in ADDED_SITES else "指定", site, queries.get(site, 0)]
            + [c[k] for k in CATEGORIES]
            + [c["_total"], c["_confirmed"], c["_ref"], " / ".join(notes.get(site, []))]
        )
    total_row = ["", "合計", sum(queries.get(s, 0) for s in counts)]
    total_row += [sum(c[k] for c in counts.values()) for k in CATEGORIES]
    total_row += [len(in_period), sum(c["_confirmed"] for c in counts.values()), len(reference), ""]
    ws.append(total_row)
    for cell in ws[1]:
        cell.fill, cell.font = header_fill, header_font
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for idx, width in enumerate([8, 28, 12, 10, 14, 10, 12, 11, 14, 12, 100], 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for r in ws.iter_rows(min_row=2):
        r[-1].alignment = Alignment(wrap_text=True, vertical="top")

    write_sheet(wb.create_sheet("一覧(2026年9月)"), in_period)
    for site in SITES:
        rows = [i for i in in_period if i["site"] == site]
        if rows:
            write_sheet(wb.create_sheet(site[:31].replace("/", "・")), rows)
    if reference:
        write_sheet(wb.create_sheet("参考(期間外)"), reference)
    wb.save(out / "new_business_2026-09.xlsx")

    print(f"items(9月)={len(in_period)} reference={len(reference)}")
    for site, c in counts.items():
        print(f"{site}\t{c['_total']}\t(確認済 {c['_confirmed']}, 参考 {c['_ref']}, queries {queries.get(site, 0)})")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
